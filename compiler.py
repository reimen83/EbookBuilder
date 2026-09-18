import io
import os
import re
import tempfile
import traceback
from PIL import Image
import pypdfium2 as pdfium
import requests

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

try:
    import docx
except ImportError:
    docx = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None


class ThemeEngine:
    CAPAS_PADRAO = {
        "music_prod": "https://images.unsplash.com/photo-1598488035139-bdbb2231ce04?q=80&w=1000",
        "finance_gold": "https://images.unsplash.com/photo-1559526324-4b87b5e36e44?q=80&w=1000",
        "ai_productivity": "https://images.unsplash.com/photo-1618005182384-a83a8bd57fbe?q=80&w=1000",
        "health_wellness": "https://images.unsplash.com/photo-1517838277536-f5f99be501cd?q=80&w=1000",
        "self_help": "https://images.unsplash.com/photo-1506126613408-eca07ce68773?q=80&w=1000",
        "dark_tech": "https://images.unsplash.com/photo-1555066931-4365d14bab8c?q=80&w=1000",
        "editorial": "https://images.unsplash.com/photo-1457369804613-52c61a468e7d?q=80&w=1000",
        "modern": "https://images.unsplash.com/photo-1486406146926-c627a92ad1ab?q=80&w=1000",
    }

    @staticmethod
    def obter_estilos(nome_tema="modern"):
        styles = getSampleStyleSheet()

        if nome_tema == "editorial":
            font_base = "Times-Roman"
            font_bold = "Times-Bold"
            cor_fundo = colors.HexColor("#FFFFFF")
            cor_primaria = colors.HexColor("#7C2D12")
            cor_titulo = colors.HexColor("#451A03")
            cor_texto = colors.HexColor("#292524")
            cor_linha = colors.HexColor("#D6D3D1")
            cor_header_footer = colors.HexColor("#78716C")
            cor_subtitulo_capa = colors.HexColor("#FDE68A")
            cor_card_capa = colors.HexColor("#1C100B")

        elif nome_tema == "dark_tech":
            font_base = "Helvetica"
            font_bold = "Helvetica-Bold"
            cor_fundo = colors.HexColor("#090D16")
            cor_primaria = colors.HexColor("#06B6D4")
            cor_titulo = colors.HexColor("#38BDF8")
            cor_texto = colors.HexColor("#E2E8F0")
            cor_linha = colors.HexColor("#0EA5E9")
            cor_header_footer = colors.HexColor("#64748B")
            cor_subtitulo_capa = colors.HexColor("#38BDF8")
            cor_card_capa = colors.HexColor("#060B13")

        elif nome_tema == "music_prod":
            font_base = "Helvetica"
            font_bold = "Helvetica-Bold"
            cor_fundo = colors.HexColor("#12131C")
            cor_primaria = colors.HexColor("#F59E0B")
            cor_titulo = colors.HexColor("#FBBF24")
            cor_texto = colors.HexColor("#F1F5F9")
            cor_linha = colors.HexColor("#8B5CF6")
            cor_header_footer = colors.HexColor("#94A3B8")
            cor_subtitulo_capa = colors.HexColor("#FBBF24")
            cor_card_capa = colors.HexColor("#0D0E15")

        elif nome_tema == "finance_gold":
            font_base = "Helvetica"
            font_bold = "Helvetica-Bold"
            cor_fundo = colors.HexColor("#FFFFFF")
            cor_primaria = colors.HexColor("#059669")
            cor_titulo = colors.HexColor("#0F172A")
            cor_texto = colors.HexColor("#1E293B")
            cor_linha = colors.HexColor("#D97706")
            cor_header_footer = colors.HexColor("#64748B")
            cor_subtitulo_capa = colors.HexColor("#6EE7B7")
            cor_card_capa = colors.HexColor("#0B1320")

        elif nome_tema == "health_wellness":
            font_base = "Helvetica"
            font_bold = "Helvetica-Bold"
            cor_fundo = colors.HexColor("#F8FAFC")
            cor_primaria = colors.HexColor("#16A34A")
            cor_titulo = colors.HexColor("#14532D")
            cor_texto = colors.HexColor("#1F2937")
            cor_linha = colors.HexColor("#86EFAC")
            cor_header_footer = colors.HexColor("#6B7280")
            cor_subtitulo_capa = colors.HexColor("#A7F3D0")
            cor_card_capa = colors.HexColor("#0A1910")

        elif nome_tema == "self_help":
            font_base = "Times-Roman"
            font_bold = "Times-Bold"
            cor_fundo = colors.HexColor("#FFFBEB")
            cor_primaria = colors.HexColor("#C2410C")
            cor_titulo = colors.HexColor("#7C2D12")
            cor_texto = colors.HexColor("#292524")
            cor_linha = colors.HexColor("#FDBA74")
            cor_header_footer = colors.HexColor("#78716C")
            cor_subtitulo_capa = colors.HexColor("#FED7AA")
            cor_card_capa = colors.HexColor("#180C07")

        elif nome_tema == "ai_productivity":
            font_base = "Helvetica"
            font_bold = "Helvetica-Bold"
            cor_fundo = colors.HexColor("#0F172A")
            cor_primaria = colors.HexColor("#7C3AED")
            cor_titulo = colors.HexColor("#38BDF8")
            cor_texto = colors.HexColor("#F1F5F9")
            cor_linha = colors.HexColor("#06B6D4")
            cor_header_footer = colors.HexColor("#94A3B8")
            cor_subtitulo_capa = colors.HexColor("#C084FC")
            cor_card_capa = colors.HexColor("#090D1A")

        else:  # 'modern'
            font_base = "Helvetica"
            font_bold = "Helvetica-Bold"
            cor_fundo = colors.HexColor("#FFFFFF")
            cor_primaria = colors.HexColor("#2563EB")
            cor_titulo = colors.HexColor("#0F172A")
            cor_texto = colors.HexColor("#334155")
            cor_linha = colors.HexColor("#2563EB")
            cor_header_footer = colors.HexColor("#64748B")
            cor_subtitulo_capa = colors.HexColor("#BFDBFE")
            cor_card_capa = colors.HexColor("#0B1528")

        h1_style = ParagraphStyle(
            "ThemeH1",
            parent=styles["Normal"],
            fontName=font_bold,
            fontSize=18,
            leading=22,
            textColor=cor_titulo,
            spaceBefore=16,
            spaceAfter=8,
        )

        h2_style = ParagraphStyle(
            "ThemeH2",
            parent=styles["Normal"],
            fontName=font_bold,
            fontSize=13,
            leading=17,
            textColor=cor_primaria,
            spaceBefore=12,
            spaceAfter=6,
        )

        body_style = ParagraphStyle(
            "ThemeBody",
            parent=styles["Normal"],
            fontName=font_base,
            fontSize=9.5,
            leading=14,
            textColor=cor_texto,
            spaceAfter=8,
        )

        bullet_style = ParagraphStyle(
            "ThemeBullet",
            parent=body_style,
            leftIndent=15,
            bulletIndent=5,
            spaceAfter=4,
        )

        table_col_style = ParagraphStyle(
            "ThemeTableCol",
            parent=body_style,
            fontSize=9,
            leading=13,
            spaceAfter=0,
        )

        return {
            "font_base": font_base,
            "font_bold": font_bold,
            "cor_fundo": cor_fundo,
            "cor_primaria": cor_primaria,
            "cor_titulo": cor_titulo,
            "cor_texto": cor_texto,
            "cor_linha": cor_linha,
            "cor_header_footer": cor_header_footer,
            "cor_subtitulo_capa": cor_subtitulo_capa,
            "cor_card_capa": cor_card_capa,
            "h1": h1_style,
            "h2": h2_style,
            "body": body_style,
            "bullet": bullet_style,
            "table_col": table_col_style,
        }


class SmartParser:
    @staticmethod
    def inferir_estrutura(texto_bruto):
        linhas = texto_bruto.split("\n")
        linhas_processadas = []

        regex_capitulo = re.compile(
            r"^(capítulo|módulo|parte|lição|introdução|conclusão|prefácio)\b",
            re.IGNORECASE,
        )
        regex_numeracao = re.compile(r"^(\d{1,2}\.|\d{1,2}\s*-\s*)")

        for linha in linhas:
            l_str = linha.strip()
            if not l_str:
                linhas_processadas.append("")
                continue

            if "|" in l_str:
                linhas_processadas.append(l_str)
                continue

            if l_str.startswith(("# ", "## ", "### ", "- ", "* ")):
                linhas_processadas.append(l_str)
                continue

            e_curto = len(l_str) <= 60
            e_caixa_alta = l_str.isupper() and len(l_str) > 3
            tem_palavra_chave = bool(regex_capitulo.match(l_str))
            sem_ponto_final = not l_str.endswith((".", ",", ";"))

            if e_curto and (tem_palavra_chave or (e_caixa_alta and sem_ponto_final)):
                linhas_processadas.append(f"# {l_str}")
                continue

            tem_num_secao = bool(regex_numeracao.match(l_str)) and len(l_str) <= 80
            termina_com_dois_pontos = l_str.endswith(":") and e_curto

            if tem_num_secao or termina_com_dois_pontos:
                linhas_processadas.append(f"## {l_str.rstrip(':')}")
                continue

            if l_str.startswith(("—", "–", "•", "1)", "2)", "3)", "a)", "b)")):
                texto_limpo = re.sub(r"^[—–•\d\w\)\-]\s*", "", l_str)
                linhas_processadas.append(f"- {texto_limpo}")
                continue

            linhas_processadas.append(l_str)

        return "\n".join(linhas_processadas)

    @staticmethod
    def extrair_titulos_documento(caminho_arquivo):
        if not caminho_arquivo or not os.path.exists(caminho_arquivo):
            return "", ""

        ext = os.path.splitext(caminho_arquivo)[1].lower()
        titulo_auto, subtitulo_auto = "", ""

        try:
            if ext in [".md", ".markdown", ".txt"]:
                with open(caminho_arquivo, "r", encoding="utf-8", errors="ignore") as f:
                    linhas = [l.strip() for l in f.readlines() if l.strip()]

                for l in linhas:
                    if not titulo_auto and (l.startswith("# ") or l.isupper()):
                        titulo_auto = re.sub(r"^#\s*", "", l)
                    elif titulo_auto and not subtitulo_auto and (l.startswith("## ") or l.startswith("### ")):
                        subtitulo_auto = re.sub(r"^#{2,3}\s*", "", l)
                        break

            elif ext == ".docx" and docx:
                doc_word = docx.Document(caminho_arquivo)
                for p in doc_word.paragraphs:
                    txt = p.text.strip()
                    if not txt:
                        continue
                    if not titulo_auto:
                        titulo_auto = txt
                    elif not subtitulo_auto:
                        subtitulo_auto = txt
                        break

            elif ext == ".pdf" and pdfplumber:
                with pdfplumber.open(caminho_arquivo) as pdf:
                    if len(pdf.pages) > 0:
                        txt_pag1 = pdf.pages[0].extract_text() or ""
                        linhas = [l.strip() for l in txt_pag1.split("\n") if l.strip()]
                        if len(linhas) > 0:
                            titulo_auto = linhas[0]
                        if len(linhas) > 1:
                            subtitulo_auto = linhas[1]

        except Exception as e:
            print(f"[Aviso Parser Titulos]: {e}")

        return titulo_auto, subtitulo_auto


class CapaHandler:
    def __init__(self, origem_capa=None, titulo="", sub_titulo="", tema="modern"):
        self.origem_capa = origem_capa
        self.titulo = titulo.strip()
        self.sub_titulo = sub_titulo.strip()
        self.tema = tema
        self.tema_config = ThemeEngine.obter_estilos(tema)

        if not self.origem_capa:
            self.origem_capa = ThemeEngine.CAPAS_PADRAO.get(
                self.tema, ThemeEngine.CAPAS_PADRAO["modern"]
            )

    def desenhar_capa(self, canvas_obj, doc):
        canvas_obj.saveState()
        largura, altura = A4
        sucesso_imagem = False

        if self.origem_capa:
            origem_str = str(self.origem_capa).strip()
            if origem_str.startswith(("http://", "https://")):
                try:
                    res = requests.get(origem_str, timeout=3)
                    if res.status_code == 200:
                        from reportlab.lib.utils import ImageReader
                        img_data = io.BytesIO(res.content)
                        img = ImageReader(img_data)
                        canvas_obj.drawImage(
                            img, 0, 0, width=largura, height=altura, preserveAspectRatio=False
                        )
                        sucesso_imagem = True
                except Exception as e:
                    print(f"[Aviso] Falha ao baixar imagem: {e}")

            elif os.path.exists(origem_str):
                try:
                    canvas_obj.drawImage(
                        origem_str, 0, 0, width=largura, height=altura, preserveAspectRatio=False
                    )
                    sucesso_imagem = True
                except Exception as e:
                    print(f"[Aviso] Falha ao carregar imagem local: {e}")

        if not sucesso_imagem:
            cor_fundo_capa = (
                self.tema_config["cor_fundo"]
                if self.tema in ["dark_tech", "music_prod", "ai_productivity"]
                else colors.HexColor("#0F172A")
            )
            canvas_obj.setFillColor(cor_fundo_capa)
            canvas_obj.rect(0, 0, largura, altura, fill=True, stroke=False)

        if self.titulo or self.sub_titulo:
            centro_x = largura / 2.0
            pos_y_bloco = altura * 0.42

            largura_card = largura * 0.84
            altura_card = 210
            x_card = (largura - largura_card) / 2.0
            y_card = pos_y_bloco - (altura_card / 2.0)

            canvas_obj.setFillColor(self.tema_config["cor_card_capa"])
            canvas_obj.setFillAlpha(0.88)
            canvas_obj.roundRect(x_card, y_card, largura_card, altura_card, 12, fill=True, stroke=False)
            canvas_obj.setFillAlpha(1.0)

            canvas_obj.setFillColor(self.tema_config["cor_primaria"])
            canvas_obj.rect(x_card + 30, y_card + altura_card - 8, largura_card - 60, 4, fill=True, stroke=False)

            if self.titulo:
                canvas_obj.setFillColor(colors.HexColor("#FFFFFF"))
                canvas_obj.setFont(self.tema_config["font_bold"], 23)
                canvas_obj.drawCentredString(centro_x, y_card + altura_card - 60, self.titulo.upper())

            canvas_obj.setFillColor(self.tema_config["cor_linha"])
            canvas_obj.rect(centro_x - 45, y_card + altura_card - 85, 90, 2, fill=True, stroke=False)

            if self.sub_titulo:
                canvas_obj.setFont(self.tema_config["font_base"], 13)
                canvas_obj.setFillColor(self.tema_config["cor_subtitulo_capa"])
                canvas_obj.drawCentredString(centro_x, y_card + 40, self.sub_titulo)

        canvas_obj.restoreState()

    def desenhar_fundo_paginas(self, canvas_obj, doc):
        if self.tema in ["dark_tech", "music_prod", "ai_productivity", "health_wellness", "self_help"]:
            canvas_obj.saveState()
            largura, altura = A4
            canvas_obj.setFillColor(self.tema_config["cor_fundo"])
            canvas_obj.rect(0, 0, largura, altura, fill=True, stroke=False)
            canvas_obj.restoreState()


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []
        self.tema_nome = "modern"

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            if self._pageNumber > 1:
                self.draw_page_decorations(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_decorations(self, page_count):
        self.saveState()
        cfg = ThemeEngine.obter_estilos(self.tema_nome)
        largura, altura = A4

        self.setFont(cfg["font_base"], 8)
        self.setFillColor(cfg["cor_header_footer"])

        self.setStrokeColor(cfg["cor_linha"])
        self.setLineWidth(0.5)
        self.line(54, 792, 541, 792)

        self.drawString(54, 36, "EbookBuilder Series")
        self.drawRightString(541, 36, f"Página {self._pageNumber} de {page_count}")
        self.line(54, 48, 541, 48)
        self.restoreState()


class EbookCompiler:
    def __init__(
        self,
        arquivo_fonte,
        arquivo_saida="ebook_final.pdf",
        capa_url=None,
        titulo_ebook="",
        sub_titulo_ebook="",
        tema="music_prod",
    ):
        self.arquivo_fonte = arquivo_fonte
        self.arquivo_saida = arquivo_saida
        self.capa_url = capa_url
        self.tema = tema
        self.theme_cfg = ThemeEngine.obter_estilos(self.tema)

        titulo_auto, subtitulo_auto = SmartParser.extrair_titulos_documento(arquivo_fonte)
        self.titulo_ebook = titulo_ebook.strip() if titulo_ebook.strip() else titulo_auto
        self.sub_titulo_ebook = sub_titulo_ebook.strip() if sub_titulo_ebook.strip() else subtitulo_auto

    def _converter_inline_formatting(self, texto):
        texto = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", texto)
        texto = re.sub(r"\*(.*?)\*", r"<i>\1</i>", texto)
        return texto

    def _parse_markdown(self, texto_md):
        story = []
        linhas = texto_md.split("\n")
        bloco_tabela = []

        for linha in linhas:
            linha_str = linha.strip()

            if "|" in linha_str:
                partes = [p.strip() for p in linha_str.split("|")]
                if len(partes) >= 2:
                    col1 = Paragraph(self._converter_inline_formatting(partes[0]), self.theme_cfg["table_col"])
                    col2 = Paragraph(self._converter_inline_formatting(partes[1]), self.theme_cfg["table_col"])
                    bloco_tabela.append([col1, col2])
                    continue

            if bloco_tabela:
                tabela = Table(bloco_tabela, colWidths=[240, 240])
                tabela.setStyle(
                    TableStyle([
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("LINEAFTER", (0, 0), (0, -1), 0.5, self.theme_cfg["cor_linha"]),
                    ])
                )
                story.append(tabela)
                story.append(Spacer(1, 8))
                bloco_tabela = []

            if not linha_str:
                continue

            if linha_str.startswith("# "):
                texto = linha_str[2:].strip()
                story.append(Paragraph(texto, self.theme_cfg["h1"]))
                story.append(
                    HRFlowable(
                        width="100%",
                        thickness=1,
                        color=self.theme_cfg["cor_linha"],
                        spaceAfter=12,
                    )
                )
            elif linha_str.startswith("## "):
                texto = linha_str[3:].strip()
                story.append(Paragraph(texto, self.theme_cfg["h2"]))
            elif linha_str.startswith("### "):
                texto = linha_str[4:].strip()
                h3_style = ParagraphStyle(
                    "CustomH3",
                    parent=self.theme_cfg["h2"],
                    fontSize=11,
                    textColor=self.theme_cfg["cor_primaria"],
                )
                story.append(Paragraph(texto, h3_style))
            elif linha_str.startswith(("- ", "* ")):
                texto = linha_str[2:].strip()
                texto_formatado = self._converter_inline_formatting(texto)
                story.append(
                    Paragraph(f"• {texto_formatado}", self.theme_cfg["bullet"])
                )
            else:
                texto_formatado = self._converter_inline_formatting(linha_str)
                story.append(Paragraph(texto_formatado, self.theme_cfg["body"]))

        if bloco_tabela:
            tabela = Table(bloco_tabela, colWidths=[240, 240])
            tabela.setStyle(
                TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("LINEAFTER", (0, 0), (0, -1), 0.5, self.theme_cfg["cor_linha"]),
                ])
            )
            story.append(tabela)

        return story

    def _parse_docx(self, caminho_docx):
        if docx is None:
            raise ImportError("A biblioteca 'python-docx' não está instalada.")
        doc_word = docx.Document(caminho_docx)
        texto_unificado = "\n".join([p.text for p in doc_word.paragraphs])
        texto_estruturado = SmartParser.inferir_estrutura(texto_unificado)
        return self._parse_markdown(texto_estruturado)

    def _parse_pdf(self, caminho_pdf):
        if pdfplumber is None:
            raise ImportError("A biblioteca 'pdfplumber' não está instalada.")
        texto_completo = []
        with pdfplumber.open(caminho_pdf) as pdf:
            for pagina in pdf.pages:
                tabelas = pagina.extract_tables()
                if tabelas:
                    for tab in tabelas:
                        for linha in tab:
                            if linha and len(linha) >= 2 and (linha[0] or linha[1]):
                                col_a = (linha[0] or "").replace("\n", " ").strip()
                                col_b = (linha[1] or "").replace("\n", " ").strip()
                                if col_a or col_b:
                                    texto_completo.append(f"{col_a} | {col_b}")
                else:
                    texto_pag = pagina.extract_text()
                    if texto_pag:
                        texto_completo.append(texto_pag)

        texto_unificado = "\n\n".join(texto_completo)
        texto_formatado = SmartParser.inferir_estrutura(texto_unificado)
        return self._parse_markdown(texto_formatado)

    def compilar(self):
        if not os.path.exists(self.arquivo_fonte):
            raise FileNotFoundError(f"Arquivo não encontrado: {self.arquivo_fonte}")

        extensao = os.path.splitext(self.arquivo_fonte)[1].lower()

        if extensao in [".md", ".markdown"]:
            with open(self.arquivo_fonte, "r", encoding="utf-8") as f:
                conteudo_flowables = self._parse_markdown(f.read())
        elif extensao == ".txt":
            with open(self.arquivo_fonte, "r", encoding="utf-8") as f:
                conteudo_flowables = self._parse_markdown(SmartParser.inferir_estrutura(f.read()))
        elif extensao == ".docx":
            conteudo_flowables = self._parse_docx(self.arquivo_fonte)
        elif extensao == ".pdf":
            conteudo_flowables = self._parse_pdf(self.arquivo_fonte)
        else:
            raise ValueError(f"Extensão não suportada: {extensao}")

        doc = SimpleDocTemplate(
            self.arquivo_saida,
            pagesize=A4,
            leftMargin=54,
            rightMargin=54,
            topMargin=72,
            bottomMargin=72,
        )

        capa_handler = CapaHandler(
            origem_capa=self.capa_url,
            titulo=self.titulo_ebook,
            sub_titulo=self.sub_titulo_ebook,
            tema=self.tema,
        )

        def criar_canvas_com_tema(*args, **kwargs):
            canvas_inst = NumberedCanvas(*args, **kwargs)
            canvas_inst.tema_nome = self.tema
            return canvas_inst

        story = [Spacer(1, 400), PageBreak()]
        story.extend(conteudo_flowables)

        doc.build(
            story,
            onFirstPage=capa_handler.desenhar_capa,
            onLaterPages=capa_handler.desenhar_fundo_paginas,
            canvasmaker=criar_canvas_com_tema,
        )

    def gerar_preview_capa_fast(self, dpi=100):
        # Utiliza o diretório temporário do SO seguro para executáveis compilados
        tmp_dir = tempfile.gettempdir()
        tmp_pdf_path = os.path.join(tmp_dir, "_capa_preview_temp.pdf")

        try:
            doc = SimpleDocTemplate(
                tmp_pdf_path,
                pagesize=A4,
                leftMargin=54,
                rightMargin=54,
                topMargin=72,
                bottomMargin=72,
            )

            capa_handler = CapaHandler(
                origem_capa=self.capa_url,
                titulo=self.titulo_ebook,
                sub_titulo=self.sub_titulo_ebook,
                tema=self.tema,
            )

            story = [Spacer(1, 400)]

            doc.build(
                story,
                onFirstPage=capa_handler.desenhar_capa,
                onLaterPages=capa_handler.desenhar_fundo_paginas,
            )

            pdf = pdfium.PdfDocument(tmp_pdf_path)
            page = pdf[0]
            image_pil = page.render(scale=dpi / 72).to_pil()
            pdf.close()

            return [image_pil]
        finally:
            if os.path.exists(tmp_pdf_path):
                try:
                    os.remove(tmp_pdf_path)
                except Exception:
                    pass
