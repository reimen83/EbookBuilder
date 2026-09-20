from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.platypus import Table
from reportlab.pdfgen import canvas

import compiler as compiler_module
from compiler import EbookCompiler, SmartParser, ThemeEngine


def test_inferir_estrutura_reconhece_titulo_secao_e_lista():
    texto = "INTRODUÇÃO\n\n1. Primeira etapa\n\n• Item importante"

    resultado = SmartParser.inferir_estrutura(texto)

    assert "# INTRODUÇÃO" in resultado
    assert "## 1. Primeira etapa" in resultado
    assert "- Item importante" in resultado


def test_extrair_titulos_de_markdown(tmp_path: Path):
    arquivo = tmp_path / "conteudo.md"
    arquivo.write_text("# Meu E-book\n\n## Um subtítulo\n\nTexto", encoding="utf-8")

    titulo, subtitulo = SmartParser.extrair_titulos_documento(str(arquivo))

    assert titulo == "Meu E-book"
    assert subtitulo == "Um subtítulo"


def test_extrair_titulos_de_texto_simples(tmp_path: Path):
    arquivo = tmp_path / "conteudo.txt"
    arquivo.write_text("Título automático\nSubtítulo automático\nTexto", encoding="utf-8")

    titulo, subtitulo = SmartParser.extrair_titulos_documento(str(arquivo))

    assert titulo == "Título automático"
    assert subtitulo == "Subtítulo automático"


def test_compilador_prioriza_titulos_manuais_e_usa_capa_do_tema(tmp_path: Path):
    arquivo = tmp_path / "conteudo.md"
    arquivo.write_text("# Título automático\n\n## Subtítulo automático", encoding="utf-8")

    compilador = EbookCompiler(
        arquivo_fonte=str(arquivo),
        titulo_ebook="Título manual",
        sub_titulo_ebook="Subtítulo manual",
        tema="modern",
    )

    assert compilador.titulo_ebook == "Título manual"
    assert compilador.sub_titulo_ebook == "Subtítulo manual"
    assert compilador.capa_url is None
    assert compilador.variacao_capa is None


def test_compilador_completa_campo_manual_com_titulo_automatico(tmp_path: Path):
    arquivo = tmp_path / "conteudo.md"
    arquivo.write_text("# Título automático\n\n## Subtítulo automático", encoding="utf-8")

    compilador = EbookCompiler(
        arquivo_fonte=str(arquivo),
        titulo_ebook="Título manual",
    )

    assert compilador.titulo_ebook == "Título manual"
    assert compilador.sub_titulo_ebook == "Subtítulo automático"


def test_capa_sem_imagem_explicita_usa_primeira_variacao_do_tema():
    from covers import CapaHandler

    capa = CapaHandler(tema="modern")
    capa_padrao = ThemeEngine.obter_opcoes_capa_por_tema("modern")[0]["url"]

    assert capa.origem_capa == capa_padrao


def test_subtitulo_longo_quebra_linhas_na_capa(tmp_path: Path):
    from covers import CapaHandler
    import pdfplumber

    capa = tmp_path / "capa.png"
    Image.new("RGB", (100, 150), color="navy").save(capa)
    saida = tmp_path / "capa-subtitulo-longo.pdf"
    pagina = canvas.Canvas(str(saida), pagesize=A4)
    CapaHandler(
        origem_capa=capa,
        titulo="Título",
        sub_titulo=(
            "Este é um subtítulo propositalmente muito longo para verificar "
            "a quebra automática dentro da caixa da capa"
        ),
        tema="modern",
    ).desenhar_capa(pagina, None)
    pagina.save()

    with pdfplumber.open(saida) as pdf:
        texto = pdf.pages[0].extract_text()
    assert texto.count("\n") >= 2
    assert "automáticadentrodacaixadacapa" in texto


def test_titulo_longo_sem_espacos_respeita_largura_da_capa():
    from covers import CapaHandler
    from reportlab.pdfbase.pdfmetrics import stringWidth

    titulo = "DVADVASDVADVSADVADVSADVADVSADVADVSADVADVSADVADVSADVADVS"
    linhas = CapaHandler._quebrar_texto(titulo, "Helvetica-Bold", 23, 368)

    assert len(linhas) > 1
    assert all(stringWidth(linha, "Helvetica-Bold", 23) <= 368 for linha in linhas)


def test_temas_têm_variacoes_e_estilos():
    opcoes = ThemeEngine.obter_opcoes_capa_por_tema("modern")
    estilos = ThemeEngine.obter_estilos("modern")

    assert opcoes
    assert all({"id", "nome", "url"} <= set(opcao) for opcao in opcoes)
    assert estilos["h1"] is not None
    assert estilos["body"] is not None


def test_compilador_rejeita_arquivo_inexistente(tmp_path: Path):
    fonte = tmp_path / "nao-existe.md"
    saida = tmp_path / "saida.pdf"

    compilador = EbookCompiler(str(fonte), str(saida))

    with pytest.raises(FileNotFoundError):
        compilador.compilar()


def test_compilador_gera_pdf_de_markdown(tmp_path: Path):
    fonte = tmp_path / "conteudo.md"
    saida = tmp_path / "saida.pdf"
    fonte.write_text("# Título\n\nUm parágrafo de teste.", encoding="utf-8")

    EbookCompiler(
        arquivo_fonte=str(fonte),
        arquivo_saida=str(saida),
        titulo_ebook="Título",
        sub_titulo_ebook="Subtítulo",
        tema="modern",
    ).compilar()

    assert saida.exists()
    assert saida.stat().st_size > 0
    assert saida.read_bytes().startswith(b"%PDF")


def test_compilador_le_docx(tmp_path: Path):
    fonte = tmp_path / "conteudo.docx"
    documento = Document()
    documento.add_paragraph("Título do documento")
    documento.add_paragraph("Texto de teste em DOCX")
    documento.save(fonte)

    compiler = EbookCompiler(str(fonte), str(tmp_path / "saida.pdf"))
    flowables = compiler._parse_docx(str(fonte))

    assert flowables


def test_compilador_le_pdf(tmp_path: Path):
    fonte = tmp_path / "conteudo.pdf"
    pdf = canvas.Canvas(str(fonte), pagesize=A4)
    pdf.drawString(72, 760, "Título do documento")
    pdf.drawString(72, 740, "Texto de teste em PDF")
    pdf.save()

    compiler = EbookCompiler(str(fonte), str(tmp_path / "saida.pdf"))
    flowables = compiler._parse_pdf(str(fonte))

    assert flowables


def test_compilador_preserva_duas_colunas_do_pdf(tmp_path: Path):
    fonte = tmp_path / "duas-colunas.pdf"
    pdf = canvas.Canvas(str(fonte), pagesize=A4)
    for indice in range(12):
        y = 780 - indice * 24
        pdf.drawString(50, y, f"Latim linha {indice}")
        pdf.drawString(310, y, f"Português linha {indice}")
    pdf.save()

    compiler = EbookCompiler(str(fonte), str(tmp_path / "saida.pdf"))
    flowables = compiler._parse_pdf(str(fonte))

    assert any(isinstance(flowable, Table) for flowable in flowables)


def test_colunas_pdf_recriam_divisorias_com_table_style(tmp_path: Path):
    fonte = tmp_path / "duas-colunas-com-linhas.pdf"
    pdf = canvas.Canvas(str(fonte), pagesize=A4)
    pdf.line(42, 700, 553, 700)
    pdf.line(297, 100, 297, 700)
    pdf.drawString(50, 680, "Coluna esquerda")
    pdf.drawString(310, 680, "Coluna direita")
    pdf.save()

    compiler = EbookCompiler(str(fonte), str(tmp_path / "saida.pdf"))
    tabela = compiler.renderer._gerar_colunas_pdf_flowable(
        [
            ("Coluna esquerda", "Coluna direita", 100),
            ("Oração esquerda", "Oração direita", 120),
        ],
        separadores=[(42, 298, 120, 297), (298, 553, 120, 297)],
        divisorias=[(297, 50, 150)],
    )

    assert tabela.horizontal_segments == [(0, 0, 1), (1, 1, 1)]
    assert not any(
        comando[0] == "LINEAFTER"
        for comando in tabela._linecmds
    )
    assert tabela.divider_segments == [(0, 0)]
    assert tabela._cellStyles[0][1].leftPadding == 8
    assert tabela._cellStyles[0][0].rightPadding == 8
    assert tabela.horizontal_segments


def test_divisorias_verticais_sao_remapeadas_ao_dividir_tabela():
    compiler = EbookCompiler("conteudo.md", "saida.pdf")
    tabela = compiler.renderer._gerar_colunas_pdf_flowable(
        [(f"Esquerda {i}", f"Direita {i}", i * 10) for i in range(8)],
        separadores=[(42, 298, 130, 297)],
        divisorias=[(297, 0, 80), (297, 80, 160)],
    )

    tabela._rowHeights = [20] * 8
    fragmentos = tabela.split(487, 70)

    assert len(fragmentos) >= 2
    assert fragmentos[0].divider_segments
    assert fragmentos[1].divider_segments
    assert fragmentos[0].divider_color is tabela.divider_color
    assert fragmentos[1].divider_color is tabela.divider_color
    assert all(
        0 <= inicio <= fim < len(fragmento._rowHeights)
        for fragmento in fragmentos
        for inicio, fim in fragmento.divider_segments
    )


def test_preview_usa_temporario_unico_e_limpa_arquivo(tmp_path: Path, monkeypatch):
    capa = tmp_path / "capa.png"
    Image.new("RGB", (40, 40), color="navy").save(capa)
    fonte = tmp_path / "conteudo.md"
    fonte.write_text("# Título", encoding="utf-8")
    caminhos_temporarios = []
    mkstemp_original = compiler_module.tempfile.mkstemp

    def mkstemp_rastreado(*args, **kwargs):
        resultado = mkstemp_original(*args, **kwargs)
        caminhos_temporarios.append(Path(resultado[1]))
        return resultado

    monkeypatch.setattr(compiler_module.tempfile, "mkstemp", mkstemp_rastreado)
    compiler = EbookCompiler(str(fonte), capa_url=str(capa))

    assert compiler.gerar_preview_capa_fast()
    assert len(caminhos_temporarios) == 1
    assert not caminhos_temporarios[0].exists()


def test_preview_usa_capa_local_e_compila_com_caminhos_path(tmp_path: Path):
    fonte = tmp_path / "conteudo com espaço.md"
    capa = tmp_path / "capa local.png"
    saida = tmp_path / "ebook final.pdf"
    fonte.write_text("# Título\n\n## Subtítulo\n\nConteúdo", encoding="utf-8")
    Image.new("RGB", (100, 150), color="navy").save(capa)

    compiler = EbookCompiler(
        arquivo_fonte=fonte,
        arquivo_saida=saida,
        capa_url=capa,
        tema="modern",
    )

    imagens = compiler.gerar_preview_capa_fast()
    compiler.compilar()

    assert imagens[0].size == (827, 1170)
    assert saida.exists()
    assert saida.read_bytes().startswith(b"%PDF")
